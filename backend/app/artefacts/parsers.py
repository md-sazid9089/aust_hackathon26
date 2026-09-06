from __future__ import annotations

import io
import re
from dataclasses import dataclass

from app.ai.guard import normalise_text

ALLOWED: dict[str, tuple[str, ...]] = {
    # extension -> acceptable declared MIME types
    "pdf": ("application/pdf",),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"),
    "txt": ("text/plain", "application/octet-stream"),
    "md": ("text/markdown", "text/plain", "application/octet-stream"),
}

CANONICAL_MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
}


class UnsupportedFile(Exception):
    pass


class NoTextExtracted(Exception):
    pass


def sniff_extension(filename: str, head: bytes) -> str:
    """Determine file type from magic bytes first, falling back to extension for plain text."""
    ext = (filename.rsplit(".", 1)[-1].lower() if "." in filename else "")
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"PK\x03\x04") and ext == "docx":
        return "docx"
    if ext in ("txt", "md"):
        if b"\x00" in head[:512]:
            raise UnsupportedFile("Binary content in a text file")
        return ext
    raise UnsupportedFile(f"Unsupported file type '{ext or 'unknown'}'")


def extract_text(ext: str, data: bytes) -> str:
    if ext == "pdf":
        text = _pdf_text(data)
    elif ext == "docx":
        text = _docx_text(data)
    else:
        text = _decode_text(data)
    text = normalise_text(text)
    if len(text) < 20:
        raise NoTextExtracted("No extractable text (scanned/image document?)")
    return text


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise UnsupportedFile("Encrypted PDF") from exc
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except PyPdfError as exc:
        raise UnsupportedFile("Corrupted or unreadable PDF") from exc


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises assorted zip/xml errors
        raise UnsupportedFile("Corrupted or unreadable DOCX") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def detect_lang(text: str) -> str:
    bn = len(re.findall(r"[\u0980-\u09FF]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    if bn == 0 and latin == 0:
        return "unknown"
    if bn == 0:
        return "en"
    if latin == 0:
        return "bn"
    return "mixed" if min(bn, latin) / max(bn, latin) > 0.15 else ("bn" if bn > latin else "en")


# ---------------------------------------------------------------------------
# Deterministic heuristic splitters (used by the mock provider and as evidence anchors)
# ---------------------------------------------------------------------------

_Q_START = re.compile(
    r"^\s*(?:Q(?:uestion)?\.?\s*)?(\d{1,2})\s*[\.\)]?\s*(?:\(?([a-h])\)?[\.\)]?)?\s+(?=\S)", re.IGNORECASE
)
_SUBPART_START = re.compile(r"^\s*\(?([a-h])\)\s+(?=\S)", re.IGNORECASE)
_MARKS = re.compile(r"[\[\(]\s*(\d{1,3}(?:\.\d+)?)\s*(?:marks?|m)?\s*[\]\)]\s*$|(\d{1,3}(?:\.\d+)?)\s*marks?\s*$", re.IGNORECASE)


@dataclass
class HeuristicQuestion:
    number: str
    text: str
    marks: float | None


def split_questions_heuristic(text: str) -> list[dict]:
    """Split exam text into numbered questions; sub-parts inherit the current main number."""
    items: list[HeuristicQuestion] = []
    current_main = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = _Q_START.match(line)
        sub = _SUBPART_START.match(line) if not m else None
        if m:
            current_main = m.group(1)
            part = m.group(2)
            number = f"{current_main}({part.lower()})" if part else current_main
            body = line[m.end():]
        elif sub and current_main:
            number = f"{current_main}({sub.group(1).lower()})"
            body = line[sub.end():]
        elif items:
            items[-1].text = f"{items[-1].text} {line}".strip()
            _pull_marks(items[-1])
            continue
        else:
            continue
        q = HeuristicQuestion(number=number, text=body.strip(), marks=None)
        _pull_marks(q)
        items.append(q)
    # a bare main number that only introduces sub-parts (e.g. "3. Answer both:") is dropped
    numbers = [q.number for q in items]
    cleaned = [
        q for q in items
        if not ("(" not in q.number and any(n.startswith(f"{q.number}(") for n in numbers) and q.marks is None)
    ]
    return [{"number": q.number, "text": q.text, "marks": q.marks if q.marks is not None else 0.0} for q in cleaned if q.text]


def _pull_marks(q: HeuristicQuestion) -> None:
    m = _MARKS.search(q.text)
    if m:
        q.marks = float(m.group(1) or m.group(2))
        q.text = q.text[: m.start()].rstrip(" -–:")


_TOPIC_LINE = re.compile(r"^\s*(?:week\s*\d+\s*[:\-.]?\s*|\d{1,2}\s*[\.\)]\s*|[-•*]\s*)(.+)$", re.IGNORECASE)


def split_topics_heuristic(text: str) -> list[dict]:
    topics: list[dict] = []
    for raw in text.splitlines():
        line = raw.strip()
        if len(line) < 4:
            continue
        m = _TOPIC_LINE.match(line)
        title = (m.group(1) if m else line).strip(" .;")
        if len(title) < 4 or len(title) > 200:
            continue
        topics.append({"code": f"T-{len(topics) + 1:02d}", "title": title})
    return topics[:60]
